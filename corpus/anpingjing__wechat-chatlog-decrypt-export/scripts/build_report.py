# -*- coding: utf-8 -*-
"""
openpyxl 生成多表 Excel 报告（含图表）
  全账号: python build_report.py --analysis analysis --out 报告.xlsx
  单群:   python build_report.py --analysis analysis --group-analysis analysis/group_xxx --out 群报告.xlsx
读取 analyze.py 产出的 summary.json 与 CSV。
"""
import argparse, csv, json, os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter


def load_csv(d, name):
    p = os.path.join(d, name)
    if not os.path.exists(p):
        return []
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.reader(f))


def coerce(rows):
    """首行表头，其余尝试数值化"""
    out = [rows[0]]
    for r in rows[1:]:
        row = []
        for v in r:
            if isinstance(v, str):
                try:
                    v = float(v) if "." in v else int(v)
                except ValueError:
                    pass
            row.append(v)
        out.append(row)
    return out


def add_sheet(wb, name, rows, widths=None, freeze=None, autofilter=None, chart=None):
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if freeze:
        ws.freeze_panes = freeze
    if autofilter:
        ws.auto_filter.ref = autofilter
    if chart and len(rows) > 1:
        kind, title = chart["type"], chart["title"]
        n = len(rows)
        ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
        ch.title = title
        ch.width, ch.height = chart.get("width", 16), chart.get("height", 9)
        data = Reference(ws, min_col=2, min_row=1, max_row=n)
        cats = Reference(ws, min_col=1, min_row=2, max_row=n)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ws.add_chart(ch, chart.get("anchor", "D2"))
    return ws


def overview_sheet(wb, summary, extra_rows=None):
    rows = [["指标", "数值"]]
    for k, v in summary.items():
        if k in ("consistency_checks", "msg_type_distribution"):
            continue
        if isinstance(v, dict):
            v = json.dumps(v, ensure_ascii=False)
        elif isinstance(v, (list, tuple)):
            v = " ~ ".join(str(x) for x in v)
        rows.append([k, v])
    if extra_rows:
        rows += extra_rows
    checks = summary.get("consistency_checks", {})
    rows.append(["一致性校验", "全部通过" if all(checks.values()) else "存在异常"])
    ws = wb.create_sheet("总览")
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 52
    return ws


def build_all(analysis_dir, out_path):
    summary = json.load(open(os.path.join(analysis_dir, "summary.json"), encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    overview_sheet(wb, summary)
    plans = [
        ("年度趋势", "年度趋势.csv", [12, 14], "A2", {"type": "bar", "title": "年度消息量"}),
        ("月度趋势", "月度趋势.csv", [12, 14], "A2", {"type": "line", "title": "月度消息量", "height": 9}),
        ("时段分布", "时段分布.csv", [16, 14], None, {"type": "pie", "title": "时段占比", "height": 10}),
        ("小时分布", "小时分布.csv", [10, 14], "A2", {"type": "bar", "title": "24小时分布"}),
        ("星期分布", "星期分布.csv", [10, 14], None, {"type": "bar", "title": "星期分布"}),
        ("消息类型", "消息类型.csv", [10, 18, 14], None, {"type": "pie", "title": "消息类型占比", "height": 10}),
        ("会话Top60", "会话Top60.csv", [38, 8, 10, 10, 12, 12, 12, 30], "A2", None),
        ("群成员文本Top80", "群成员文本Top80.csv", [38, 16, 14], "A2", None),
    ]
    for name, csv_name, widths, freeze, chart in plans:
        rows = coerce(load_csv(analysis_dir, csv_name))
        if rows:
            add_sheet(wb, name, rows, widths, freeze,
                      f"A1:{get_column_letter(len(rows[0]))}{len(rows)}" if freeze else None, chart)
    wb.save(out_path)
    print("saved:", out_path)


def build_group(group_dir, out_path):
    summary = json.load(open(os.path.join(group_dir, "summary.json"), encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    top_day = (summary.get("最活跃5天") or [["-", 0]])[0]
    overview_sheet(wb, summary, extra_rows=[
        ["最活跃一天", f'{top_day[0]}（{top_day[1]}条）'],
    ])
    plans = [
        ("成员发言Top50", "群_成员发言Top50.csv", [18, 12, 10, 14, 14], "A2", None),
        ("月度趋势", "群_月度.csv", [12, 12], None, {"type": "line", "title": "月度消息量", "height": 9}),
        ("年度趋势", "群_年度.csv", [10, 12], None, {"type": "bar", "title": "年度消息量"}),
        ("小时分布", "群_小时.csv", [10, 12], None, {"type": "bar", "title": "24小时分布"}),
        ("星期分布", "群_星期.csv", [10, 12], None, {"type": "bar", "title": "星期分布"}),
        ("消息类型", "群_类型.csv", [10, 16, 12], None, None),
        ("活跃日Top30", "群_活跃日Top30.csv", [14, 12], "A2", None),
    ]
    for name, csv_name, widths, freeze, chart in plans:
        rows = coerce(load_csv(group_dir, csv_name))
        if rows:
            add_sheet(wb, name, rows, widths, freeze,
                      f"A1:{get_column_letter(len(rows[0]))}{len(rows)}" if freeze else None, chart)
    wb.save(out_path)
    print("saved:", out_path)


def main():
    ap = argparse.ArgumentParser(description="生成 Excel 分析报告")
    ap.add_argument("--analysis", help="analyze.py 全账号输出目录")
    ap.add_argument("--group-analysis", help="analyze.py --group 输出目录")
    ap.add_argument("--out", required=True, help="输出 xlsx 路径")
    args = ap.parse_args()
    if args.group_analysis:
        build_group(args.group_analysis, args.out)
    elif args.analysis:
        build_all(args.analysis, args.out)
    else:
        ap.error("需要 --analysis 或 --group-analysis 之一")


if __name__ == "__main__":
    main()
