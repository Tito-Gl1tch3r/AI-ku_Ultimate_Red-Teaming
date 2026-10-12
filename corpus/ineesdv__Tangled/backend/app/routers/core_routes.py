from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
import base64
from sqlalchemy.orm import Session
from sqlalchemy import func, exists
from typing import List
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo
from ..database import get_db
from ..models import *
from ..schemas import *
from ..auth import get_current_user, get_user_from_api_key, get_user_both_auths
from cryptography.fernet import Fernet
from ..config import settings
import json
import re
from typing import Set
from sqlalchemy import select, exists, text
from fastapi import Body
from sqlalchemy.exc import IntegrityError
from fastapi.responses import StreamingResponse
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

router = APIRouter(prefix="/api")

DATETIME_RE = re.compile(r'^\d{8}T\d{6}$')
VAR_TOKEN_RE = re.compile(r"\{\{\s*([A-Za-z0-9_.:-]+)\s*(?:\|[^}]*)?\}\}")

def _validate_dt(s: str | None) -> None:
    if s and not DATETIME_RE.match(s):
        raise HTTPException(status_code=422, detail="meeting_* must be YYYYMMDDTHHMMSS")

def _iso_utc(dt: datetime | None):
    if not dt:
        return None
    
    if dt.tzinfo is None:
        # stored as UTC
        return dt.isoformat() + "Z"
    
    return dt.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")

def _fernet() -> Fernet:
    if not settings.fernet_key:
        raise HTTPException(status_code=500, detail="FERNET_KEY not configured")
    
    return Fernet(settings.fernet_key.encode() if isinstance(settings.fernet_key, str) else settings.fernet_key)

def _snapshot_calendar_fields_if_empty(db: Session, c: Campaign, template: EmailTemplate):
    if not template:
        return

    # Snapshot calendar mode
    if getattr(c, "is_calendar_cancellation", None) is None:
        c.is_calendar_cancellation = bool(getattr(template, "is_calendar_cancellation", False))
    if getattr(c, "is_calendar_reschedule", None) is None:
        c.is_calendar_reschedule = bool(getattr(template, "is_calendar_reschedule", False))

    # Snapshot relevant fields
    if template.calendar_enabled or template.is_calendar_cancellation or template.is_calendar_reschedule:
        changed = False

        # Common fields
        if not c.meeting_summary and template.meeting_summary:
            c.meeting_summary = template.meeting_summary; changed = True
        if not c.meeting_begin and template.meeting_begin:
            c.meeting_begin = template.meeting_begin; changed = True
        if not c.meeting_end and template.meeting_end:
            c.meeting_end = template.meeting_end; changed = True
        if not c.meeting_tz and template.meeting_tz:
            c.meeting_tz = template.meeting_tz; changed = True

        # UID & Sequence
        if not c.calendar_uid and template.meeting_uid:
            c.calendar_uid = template.meeting_uid; changed = True

        if getattr(c, "meeting_sequence", None) is None and getattr(template, "meeting_sequence", None) is not None:
            c.meeting_sequence = template.meeting_sequence; changed = True

        # Provider & URL snapshot (don't snapshot in cancellations)
        if (template.calendar_enabled and not template.is_calendar_cancellation) or template.is_calendar_reschedule:
            if (not getattr(c, "meeting_provider", None)) or c.meeting_provider == "none":
                c.meeting_provider = (template.meeting_provider or "none"); changed = True
            if (not getattr(c, "meeting_url", None)) and template.meeting_provider in ("teams", "gmeet") and template.meeting_url:
                c.meeting_url = template.meeting_url; changed = True
        else:
            # If cancellation, ensure meeting url is empty
            if getattr(c, "meeting_provider", None) not in (None, "none"):
                c.meeting_provider = "none"; changed = True
            if getattr(c, "meeting_url", None):
                c.meeting_url = None; changed = True

        # Include all targets
        if getattr(c, "include_all_targets_as_attendees", None) is None:
            if template.is_calendar_cancellation:
                c.include_all_targets_as_attendees = False
            else:
                c.include_all_targets_as_attendees = bool(getattr(template, "include_all_targets_as_attendees", False))

        if changed:
            db.commit()

def _extract_vars_from_text(text: str | None) -> set[str]:
    if not text:
        return set()
    
    return set(m.group(1) for m in VAR_TOKEN_RE.finditer(text))

def _scan_template_subject_body(template) -> set[str]:
    used: set[str] = set()

    for attr in ("subject", "body_html", "body_text"):
        used |= _extract_vars_from_text(getattr(template, attr, None))
    
    return used

def _scan_calendar_meeting_from_campaign(campaign, template) -> set[str]:
    used: set[str] = set()

    if getattr(template, "calendar_enabled", False) or getattr(template, "is_calendar_cancellation", False):
        # Calendar text fields
        for attr in ("meeting_summary", "location"):
            used |= _extract_vars_from_text(getattr(campaign, attr, None))

        # Meeting URL only for invites/reschedule and if provider is set
        if getattr(template, "calendar_enabled", False) or getattr(template, "is_calendar_reschedule", False):
            prov = getattr(campaign, "meeting_provider", None) or getattr(template, "meeting_provider", None)
            if prov and prov != "none":
                url = getattr(campaign, "meeting_url", None) or getattr(template, "meeting_url", None)
                used |= _extract_vars_from_text(url)
    
    return used

def _snapshot_used_vars_if_empty(db, campaign, template):
    if getattr(campaign, "used_var_names", None) is not None:
        return
    
    candidates = _scan_template_subject_body(template) | _scan_calendar_meeting_from_campaign(campaign, template)
    campaign.used_var_names = sorted(list(candidates))
    db.commit()

def _track_used_vars_from_log(db, campaign, template, gen_vars: dict | None):
    if not gen_vars:
        return
    
    current = set(campaign.used_var_names or [])
    candidates = set(current) or (_scan_template_subject_body(template) | _scan_calendar_meeting_from_campaign(campaign, template))
    actually_used = set(gen_vars.keys()) & candidates

    if not actually_used:
        return
    
    newset = current | actually_used

    if newset != current:
        campaign.used_var_names = sorted(list(newset))
        db.commit()

ALWAYS_SHOW = {"name", "email", "extraField"}

def _campaign_used_vars(db, campaign, template) -> set[str]:
    # Always include basic fields
    base = set(ALWAYS_SHOW)

    # If campaign has started return snapshotted values
    snap = getattr(campaign, "used_var_names", None)
    if snap is not None and getattr(campaign, "started_at", None):
        return base | set(snap)

    # If it hasn't started scan template & campaign info
    started = bool(getattr(campaign, "started_at", None))
    if not started:
        return base | _scan_template_subject_body(template) | _scan_calendar_meeting_from_campaign(campaign, template)

    # Started but no snapshot -> Scan logs
    used: set[str] = set()
    for lg in db.query(EmailLog).filter(EmailLog.campaign_id == campaign.id):
        if getattr(lg, "gen_vars", None):
            used.update(lg.gen_vars.keys())

    used |= _scan_calendar_meeting_from_campaign(campaign, template)
    return base | used

def _cleanup_pending_emails(db: Session, campaign_id: int, new_status: str = "cancelled"):
    updated = (
        db.query(EmailLog)
        .filter(
            EmailLog.campaign_id == campaign_id,
            EmailLog.status.in_(["pending", "retrying"])
        )
        .update(
            {
                EmailLog.status: new_status,
                EmailLog.error: f"Campaign ended without sending this email"
            },
            synchronize_session=False
        )
    )
    if updated > 0:
        print(f"[campaign {campaign_id}] Cleaned up {updated} pending/retrying emails to status '{new_status}'", flush=True)
    
    return updated

@router.get("/stats", response_model=StatsOut, tags=["Campaigns"])
def stats(db: Session = Depends(get_db), user=Depends(get_current_user)):
    total_campaigns = db.query(Campaign).count()
    active_campaigns = db.query(Campaign).filter(Campaign.status=="running").count()
    paused_campaigns = db.query(Campaign).filter(Campaign.status=="paused").count()
    archived_campaigns = db.query(Campaign).filter(Campaign.status=="archived").count()
    scheduled_campaigns = db.query(Campaign).filter(Campaign.status=="scheduled").count()

    total_targets = db.query(Target).count()
    today = date.today()
    emails_sent_today = db.query(EmailLog).filter(EmailLog.status=="sent", func.date(EmailLog.sent_at)==today).count()

    return StatsOut(
        total_campaigns=total_campaigns,
        active_campaigns=active_campaigns,
        paused_campaigns=paused_campaigns,
        archived_campaigns=archived_campaigns,
        scheduled_campaigns=scheduled_campaigns,
        total_targets=total_targets,
        emails_sent_today=emails_sent_today
    )

@router.post("/servers", response_model=ServerOut, tags=["Servers"])
def create_server(payload: ServerCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if db.query(exists().where(SendingServer.name == payload.name)).scalar():
        raise HTTPException(status_code=409, detail="A sending server with that name already exists")
    
    s = SendingServer(name=payload.name, host=payload.host, port=payload.port, use_tls=payload.use_tls)
    db.add(s)

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A sending server with that name already exists")
    
    db.refresh(s)
    return ServerOut(id=s.id, name=s.name, host=s.host, port=s.port, use_tls=s.use_tls)

@router.get("/servers", response_model=List[ServerOut], tags=["Servers"])
def list_servers(db: Session = Depends(get_db), user=Depends(get_current_user)):
    return [ServerOut(id=s.id, name=s.name, host=s.host, port=s.port, use_tls=s.use_tls) for s in db.query(SendingServer).order_by(SendingServer.name.asc()).all()]

@router.put("/servers/{server_id}", response_model=ServerOut, tags=["Servers"])
def update_server(server_id: int, payload: ServerUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    s = db.query(SendingServer).get(server_id)
    
    if not s:
        raise HTTPException(status_code=404, detail="Not found")

    if payload.name is not None and payload.name != s.name:
        if db.query(exists().where(SendingServer.name == payload.name)).scalar():
            raise HTTPException(status_code=409, detail="A sending server with that name already exists")
    
    for f in ["name","host","port","use_tls"]:
        v = getattr(payload, f)
        if v is not None: setattr(s, f, v)
    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A sending server with that name already exists")

    db.refresh(s)
    return ServerOut(id=s.id, name=s.name, host=s.host, port=s.port, use_tls=s.use_tls)

@router.delete("/servers/{server_id}", tags=["Servers"])
def delete_server(server_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    s = db.query(SendingServer).get(server_id)
    
    if not s:
        raise HTTPException(status_code=404, detail="Not found")

    # Block deletion if any active campaign still depends on it
    active_statuses = ("running", "paused", "scheduled")
    active = db.query(Campaign).filter(
        Campaign.server_id == s.id,
        Campaign.status.in_(active_statuses)
    ).all()

    if active:
        names = ", ".join(c.name for c in active)
        raise HTTPException(
            status_code=409,
            detail=f"Server is used by active campaign(s): {names}. Stop/archive them before deleting."
        )

    # Snapshot missing
    db.query(Campaign).filter(
        Campaign.server_id == s.id,
        Campaign.server_name_snapshot.is_(None),
    ).update(
        {Campaign.server_name_snapshot: s.name},
        synchronize_session=False,
    )

    db.delete(s)
    db.commit()
    return {"ok": True}

@router.post("/servers/{server_id}/duplicate", response_model=ServerOut, tags=["Servers"])
def duplicate_server(server_id: int, payload: ServerDuplicate = Body(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    src = db.query(SendingServer).filter(SendingServer.id == server_id).first()

    if not src:
        raise HTTPException(status_code=404, detail="Server not found")

    exists_ = db.query(exists().where(SendingServer.name == payload.name)).scalar()
    if exists_:
        raise HTTPException(status_code=409, detail="A sending server with that name already exists")

    dst = SendingServer(
        name=payload.name,
        host=src.host,
        port=src.port,
        use_tls=src.use_tls
    )

    db.add(dst)
    db.commit()
    db.refresh(dst)

    return ServerOut(id=dst.id, name=dst.name, host=dst.host, port=dst.port, use_tls=dst.use_tls)

@router.post("/profiles", response_model=ProfileOut, tags=["Sender profiles"])
def create_profile(payload: ProfileCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if db.query(exists().where(SenderProfile.profile_name == payload.profile_name)).scalar():
        raise HTTPException(status_code=409, detail="A sender profile with that name already exists")
    
    enc = _fernet().encrypt(payload.password.encode())
    p = SenderProfile(profile_name=payload.profile_name, display_name=payload.display_name, email=payload.email, enc_password=enc)
    db.add(p)

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A sender profile with that name already exists")
    
    db.refresh(p)

    return ProfileOut(id=p.id, profile_name=p.profile_name, display_name=p.display_name, email=p.email)

@router.get("/profiles", response_model=List[ProfileOut], tags=["Sender profiles"])
def list_profiles(db: Session = Depends(get_db), user=Depends(get_current_user)):
    return [ProfileOut(id=p.id, profile_name=p.profile_name, display_name=p.display_name, email=p.email) for p in db.query(SenderProfile).order_by(SenderProfile.profile_name.asc()).all()]

@router.put("/profiles/{profile_id}", response_model=ProfileOut, tags=["Sender profiles"])
def update_profile(profile_id: int, payload: ProfileUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = db.query(SenderProfile).get(profile_id)

    if not p:
        raise HTTPException(status_code=404, detail="Not found")
    
    if payload.profile_name is not None and payload.profile_name != p.profile_name:
        if db.query(exists().where(SenderProfile.profile_name == payload.profile_name)).scalar():
            raise HTTPException(status_code=409, detail="A sender profile with that name already exists")
        
    for f in ["profile_name","display_name","email"]:
        v = getattr(payload, f)
        if v is not None:
            setattr(p, f, v)
    
    if payload.password:
        p.enc_password = _fernet().encrypt(payload.password.encode())
    
    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A sender profile with that profile_name already exists")

    db.refresh(p)

    return ProfileOut(id=p.id, profile_name=p.profile_name, display_name=p.display_name, email=p.email)

@router.delete("/profiles/{profile_id}", tags=["Sender profiles"])
def delete_profile(profile_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = db.query(SenderProfile).get(profile_id)

    if not p:
        raise HTTPException(status_code=404, detail="Not found")

    active_statuses = ("running", "paused", "scheduled")

    active = db.query(Campaign).filter(
        Campaign.profile_id == p.id,
        Campaign.status.in_(active_statuses)
    ).all()

    if active:
        names = ", ".join(c.name for c in active)
        raise HTTPException(
            status_code=409,
            detail=f"Sender profile is used by active campaign(s): {names}. Stop/archive them before deleting."
        )

    # Snapshot missing
    db.query(Campaign).filter(
        Campaign.profile_id == p.id,
        Campaign.profile_name_snapshot.is_(None),
    ).update(
        {
            Campaign.profile_name_snapshot: p.profile_name,
            Campaign.profile_email_snapshot: p.email,
        },
        synchronize_session=False,
    )

    db.delete(p)
    db.commit()
    return {"ok": True}

@router.post("/profiles/{profile_id}/duplicate", response_model=ProfileOut, tags=["Sender profiles"])
def duplicate_profile(profile_id: int, payload: ProfileDuplicate = Body(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    src = db.query(SenderProfile).filter(SenderProfile.id == profile_id).first()

    if not src:
        raise HTTPException(status_code=404, detail="Profile not found")

    exists_ = db.query(exists().where(SenderProfile.profile_name == payload.profile_name)).scalar()
    if exists_:
        raise HTTPException(status_code=409, detail="A sender profile with that name already exists")

    dst = SenderProfile(
        profile_name=payload.profile_name,
        display_name=src.display_name,
        email=src.email,
        enc_password=src.enc_password, # copy encrypted password
    )

    db.add(dst)
    db.commit()
    db.refresh(dst)

    return ProfileOut(id=dst.id, profile_name=dst.profile_name, display_name=dst.display_name, email=dst.email)

@router.post("/templates", response_model=TemplateOut, tags=["Templates"])
def create_template(payload: TemplateCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if db.query(exists().where(EmailTemplate.name == payload.name)).scalar():
        raise HTTPException(status_code=409, detail="A template with that name already exists")
    
    # Exclusivity validation
    if (payload.is_calendar_cancellation and payload.is_calendar_reschedule):
        raise HTTPException(status_code=422, detail="A template cannot be both a reschedule and a cancellation")

    # Validations for invites
    if payload.calendar_enabled and not payload.is_calendar_cancellation and not payload.is_calendar_reschedule:
        _validate_dt(payload.meeting_begin)
        _validate_dt(payload.meeting_end)

        if not payload.organizer_email:
            raise HTTPException(status_code=422, detail="organizer_email is required when calendar_enabled")
        
        # Ignore any UID/sequence provided by the client
        payload.meeting_uid = None
        payload.meeting_sequence = None

    # Validations for reschedules
    elif payload.is_calendar_reschedule:
        _validate_dt(payload.meeting_begin)
        _validate_dt(payload.meeting_end)

        if not payload.organizer_email:
            raise HTTPException(status_code=422, detail="organizer_email is required for reschedule")
        
        if not payload.meeting_uid:
            raise HTTPException(status_code=422, detail="meeting_uid is required for reschedule")
        
        if payload.meeting_sequence is None or not isinstance(payload.meeting_sequence, int) or payload.meeting_sequence < 0:
            raise HTTPException(status_code=422, detail="meeting_sequence (>= 0) is required for reschedule")

    # Validations for cancellations
    elif payload.is_calendar_cancellation:
        _validate_dt(payload.meeting_begin)
        _validate_dt(payload.meeting_end)

        if not payload.organizer_email:
            raise HTTPException(status_code=422, detail="organizer_email is required for calendar cancellation")
        
        if not payload.meeting_uid:
            raise HTTPException(status_code=422, detail="meeting_uid is required for calendar cancellation")
        
        if payload.meeting_sequence is None or not isinstance(payload.meeting_sequence, int) or payload.meeting_sequence < 0:
            raise HTTPException(status_code=422, detail="meeting_sequence (>= 0) is required for calendar cancellation")

        # Exclude meeting urls
        payload.meeting_provider = "none"
        payload.meeting_url = None

        # Force not including all attendees
        payload.include_all_targets_as_attendees = False

    t = EmailTemplate(
        name=payload.name,
        subject=payload.subject,
        body_html=payload.body_html,
        body_text=payload.body_text,
        calendar_enabled=payload.calendar_enabled,
        meeting_summary=payload.meeting_summary,
        organizer_cn=payload.organizer_cn,
        organizer_email=str(payload.organizer_email) if payload.organizer_email else None,
        attendees=[a.dict() for a in (payload.attendees or [])],
        disable_forwarding=payload.disable_forwarding,
        calendar_priority=payload.calendar_priority or 2,
        calendar_categories = (payload.calendar_categories or None),
        calendar_private=payload.calendar_private,
        meeting_begin=payload.meeting_begin,
        meeting_end=payload.meeting_end,
        meeting_provider=(payload.meeting_provider or "none"),
        meeting_url=payload.meeting_url,
        meeting_tz=payload.meeting_tz,
        location=payload.location,
        allow_responses=payload.allow_responses,
        include_all_targets_as_attendees=bool(getattr(payload, "include_all_targets_as_attendees", False)),
        is_calendar_reschedule=bool(payload.is_calendar_reschedule),
        is_calendar_cancellation=bool(payload.is_calendar_cancellation),
        meeting_uid=payload.meeting_uid,
        meeting_sequence=payload.meeting_sequence,
    )
    
    db.add(t)

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A template with that name already exists")

    db.refresh(t)

    if payload.attachments:
        for a in payload.attachments:
            raw_b64 = (a.get("base64") or "")
            raw = base64.b64decode(raw_b64.split(",")[-1]) if raw_b64 else b""
            att = TemplateAttachment(
                template_id=t.id,
                filename=a.get("filename", "file"),
                mime_type=a.get("mime_type", "application/octet-stream"),
                data=raw,
            )
            db.add(att)
        db.commit()

    atts = db.query(TemplateAttachment).filter(TemplateAttachment.template_id == t.id).all()
    att_out = [AttachmentOut(id=a.id, filename=a.filename, mime_type=a.mime_type, size=len(a.data or b"")) for a in atts]

    attachments_count = db.query(func.count(TemplateAttachment.id))\
        .filter(TemplateAttachment.template_id == t.id).scalar() or 0

    calendar_attachments_count = db.query(func.count(TemplateCalendarAttachment.id))\
        .filter(TemplateCalendarAttachment.template_id == t.id).scalar() or 0

    raw_attendees = getattr(t, "attendees", None)
    if isinstance(raw_attendees, str):
        try:
            attendees = json.loads(raw_attendees) or []
        except Exception:
            attendees = []
    elif isinstance(raw_attendees, list):
        attendees = raw_attendees
    else:
        attendees = []

    return TemplateOut(
        id=t.id,
        name=t.name,
        subject=t.subject,
        body_html=t.body_html,
        body_text=t.body_text,
        attachments=att_out,
        attachments_count=attachments_count,
        calendar_attachments_count=calendar_attachments_count,
        calendar_enabled=bool(getattr(t, "calendar_enabled", False) or False),
        meeting_summary=getattr(t, "meeting_summary", None),
        organizer_cn=getattr(t, "organizer_cn", None),
        organizer_email=getattr(t, "organizer_email", None),
        attendees=attendees,
        disable_forwarding=bool(getattr(t, "disable_forwarding", False) or False),
        calendar_priority=int(getattr(t, "calendar_priority", 2) or 2),
        calendar_categories=(getattr(t, "calendar_categories", None) or None),
        calendar_private=getattr(t, "calendar_private", None),
        meeting_begin=getattr(t, "meeting_begin", None),
        meeting_end=getattr(t, "meeting_end", None),
        meeting_provider=(getattr(t, "meeting_provider", "none") or "none"),
        meeting_url=getattr(t, "meeting_url", None),
        meeting_tz=getattr(t, "meeting_tz", None),
        location=getattr(t, "location", None),
        allow_responses=bool(getattr(t, "allow_responses", True)),
        include_all_targets_as_attendees=bool(getattr(t, "include_all_targets_as_attendees", False)),
        is_calendar_reschedule=bool(getattr(t, "is_calendar_reschedule", False)),
        is_calendar_cancellation=bool(getattr(t, "is_calendar_cancellation", False)),
        meeting_uid=getattr(t, "meeting_uid", None),
        meeting_sequence=getattr(t, "meeting_sequence", None),
    )

@router.get("/templates", response_model=List[TemplateOut], tags=["Templates"])
def list_templates(db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(EmailTemplate).order_by(EmailTemplate.name.asc()).all()
    out: List[TemplateOut] = []
    
    # Attachments
    for t in rows:
        atts = (
            db.query(TemplateAttachment)
              .filter(TemplateAttachment.template_id == t.id)
              .all()
        )
        att_out = [
            AttachmentOut(
                id=a.id,
                filename=a.filename,
                mime_type=a.mime_type,
                size=len(a.data or b""),
            ) for a in atts
        ]

        attachments_count = db.query(func.count(TemplateAttachment.id))\
            .filter(TemplateAttachment.template_id == t.id).scalar() or 0

        calendar_attachments_count = db.query(func.count(TemplateCalendarAttachment.id))\
            .filter(TemplateCalendarAttachment.template_id == t.id).scalar() or 0

        # Attendees
        raw_attendees = getattr(t, "attendees", None)
        if isinstance(raw_attendees, str):
            try:
                attendees = json.loads(raw_attendees) or []
            except Exception:
                attendees = []
        elif isinstance(raw_attendees, list):
            attendees = raw_attendees
        elif raw_attendees is None:
            attendees = []
        else:
            attendees = list(raw_attendees) if hasattr(raw_attendees, "__iter__") else []

        # Redundant check but just in case - no meeting urls in cancellations
        meeting_provider = (getattr(t, "meeting_provider", "none") or "none")
        meeting_url = getattr(t, "meeting_url", None)
        if getattr(t, "is_calendar_cancellation", False):
            meeting_provider = "none"
            meeting_url = None

        out.append(TemplateOut(
            id=t.id,
            name=t.name,
            subject=t.subject,
            body_html=t.body_html,
            body_text=t.body_text,
            attachments=att_out,
            attachments_count=attachments_count,
            calendar_attachments_count=calendar_attachments_count,
            calendar_enabled=bool(getattr(t, "calendar_enabled", False) or False),
            meeting_summary=getattr(t, "meeting_summary", None),
            organizer_cn=getattr(t, "organizer_cn", None),
            organizer_email=getattr(t, "organizer_email", None),
            attendees=attendees,
            disable_forwarding=bool(getattr(t, "disable_forwarding", False) or False),
            calendar_priority=int(getattr(t, "calendar_priority", 2) or 2),
            calendar_categories=(getattr(t, "calendar_categories", None) or None),
            calendar_private=getattr(t, "calendar_private", None),
            meeting_begin=getattr(t, "meeting_begin", None),
            meeting_end=getattr(t, "meeting_end", None),
            meeting_provider=meeting_provider,
            meeting_url=meeting_url,
            meeting_tz=getattr(t, "meeting_tz", None),
            location=getattr(t, "location", None),
            allow_responses=bool(getattr(t, "allow_responses", True)),
            include_all_targets_as_attendees=bool(getattr(t, "include_all_targets_as_attendees", False)),
            is_calendar_reschedule=bool(getattr(t, "is_calendar_reschedule", False)),
            is_calendar_cancellation=bool(getattr(t, "is_calendar_cancellation", False)),
            meeting_uid=getattr(t, "meeting_uid", None),
            meeting_sequence=getattr(t, "meeting_sequence", None),
        ))
    return out

def was_sent(model, field: str) -> bool:
    return field in getattr(model, "model_fields_set", set())

@router.put("/templates/{template_id}", response_model=TemplateOut, tags=["Templates"])
def update_template(template_id: int, payload: TemplateUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")

    if payload.name is not None and payload.name != t.name:
        if db.query(exists().where(EmailTemplate.name == payload.name)).scalar():
            raise HTTPException(status_code=409, detail="A template with that name already exists")

    if was_sent(payload, "name"):
        t.name = payload.name
    if was_sent(payload, "subject"):
        t.subject = payload.subject
    if was_sent(payload, "body_html"):
        t.body_html = payload.body_html
    if was_sent(payload, "body_text"):
        t.body_text = payload.body_text

    if payload.calendar_enabled is not None:
        t.calendar_enabled = payload.calendar_enabled

    # Exclusivity
    flag_cancel = getattr(payload, "is_calendar_cancellation", None)
    flag_resched = getattr(payload, "is_calendar_reschedule", None)

    if flag_cancel is not None:
        t.is_calendar_cancellation = bool(flag_cancel)

    if flag_resched is not None:
        t.is_calendar_reschedule = bool(flag_resched)

    if t.is_calendar_cancellation and t.calendar_enabled:
        raise HTTPException(status_code=422, detail="Template cannot be both invite and cancellation")

    if t.is_calendar_reschedule and t.calendar_enabled:
        raise HTTPException(status_code=422, detail="Template cannot be both invite and reschedule")

    if t.is_calendar_cancellation and t.is_calendar_reschedule:
        raise HTTPException(status_code=422, detail="Template cannot be both reschedule and cancellation")

    # No invite when cancel/reschedule
    if t.is_calendar_cancellation or t.is_calendar_reschedule:
        t.calendar_enabled = False

    # Normalize fields 
    if t.is_calendar_cancellation:
        # No conferencing
        t.meeting_provider = "none"
        t.meeting_url = None

        # Delete all attached files
        db.query(TemplateCalendarAttachment).filter(
            TemplateCalendarAttachment.template_id == t.id
        ).delete(synchronize_session=False)

    elif t.is_calendar_reschedule:
        pass

    elif t.calendar_enabled:
        t.meeting_uid = None
        t.meeting_sequence = None

    # Calendar fields
    if payload.meeting_summary is not None:
        t.meeting_summary = payload.meeting_summary
    if payload.organizer_cn is not None:
        t.organizer_cn = payload.organizer_cn
    if payload.organizer_email is not None:
        t.organizer_email = str(payload.organizer_email)
    if payload.attendees is not None:
        t.attendees = [a.dict() for a in payload.attendees]
    if payload.disable_forwarding is not None:
        t.disable_forwarding = payload.disable_forwarding
    if payload.calendar_priority is not None:
        t.calendar_priority = payload.calendar_priority
    if was_sent(payload, "calendar_categories"):
        t.calendar_categories = payload.calendar_categories
    if was_sent(payload, "calendar_private"):
        t.calendar_private = payload.calendar_private
    if payload.meeting_begin is not None:
        _validate_dt(payload.meeting_begin)
        t.meeting_begin = payload.meeting_begin
    if payload.meeting_end is not None:
        _validate_dt(payload.meeting_end)
        t.meeting_end = payload.meeting_end
    if t.is_calendar_cancellation:
        # Force not including all attendees
        t.include_all_targets_as_attendees = False
    else:
        if was_sent(payload, "include_all_targets_as_attendees"):
            t.include_all_targets_as_attendees = bool(payload.include_all_targets_as_attendees)
    
    # Conferencing (skip if cancellation)
    if not bool(getattr(t, "is_calendar_cancellation", False)):
        if getattr(payload, "meeting_provider", None) is not None:
            t.meeting_provider = payload.meeting_provider or "none"
            if t.meeting_provider == "none":
                t.meeting_url = None
        if payload.meeting_url is not None:
            t.meeting_url = payload.meeting_url

    if payload.meeting_tz is not None:
        t.meeting_tz = payload.meeting_tz
    if was_sent(payload, "location"):
        t.location = payload.location
    if payload.allow_responses is not None:
        t.allow_responses = payload.allow_responses
    if payload.meeting_uid is not None:
        t.meeting_uid = payload.meeting_uid
    if payload.meeting_sequence is not None:
        t.meeting_sequence = payload.meeting_sequence

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A template with that name already exists")

    db.refresh(t)

    if payload.attachments is not None:
        # Drop old ones
        db.query(TemplateAttachment).filter(TemplateAttachment.template_id == t.id).delete()
        db.commit()

        # Add the new set
        for a in payload.attachments or []:
            raw_b64 = (a.get("base64") or "")
            raw = base64.b64decode(raw_b64.split(",")[-1]) if raw_b64 else b""
            att = TemplateAttachment(
                template_id=t.id,
                filename=a.get("filename", "file"),
                mime_type=a.get("mime_type", "application/octet-stream"),
                data=raw,
            )
            db.add(att)
        db.commit()

    atts = db.query(TemplateAttachment).filter(TemplateAttachment.template_id == t.id).all()
    att_out = [AttachmentOut(id=a.id, filename=a.filename, mime_type=a.mime_type, size=len(a.data or b"")) for a in atts]

    attachments_count = db.query(func.count(TemplateAttachment.id)).filter(TemplateAttachment.template_id == t.id).scalar() or 0

    calendar_attachments_count = db.query(func.count(TemplateCalendarAttachment.id)).filter(TemplateCalendarAttachment.template_id == t.id).scalar() or 0

    # Normalize for cancellation
    if getattr(t, "is_calendar_cancellation", False):
        calendar_attachments_count = 0

    # Attendees
    raw_attendees = getattr(t, "attendees", None)
    if isinstance(raw_attendees, str):
        try:
            attendees = json.loads(raw_attendees) or []
        except Exception:
            attendees = []
    elif isinstance(raw_attendees, list):
        attendees = raw_attendees
    else:
        attendees = []

    return TemplateOut(
        id=t.id,
        name=t.name,
        subject=t.subject,
        body_html=t.body_html,
        body_text=t.body_text,
        attachments=att_out,
        attachments_count=attachments_count,
        calendar_attachments_count=calendar_attachments_count,
        calendar_enabled=bool(getattr(t, "calendar_enabled", False) or False),
        meeting_summary=getattr(t, "meeting_summary", None),
        organizer_cn=getattr(t, "organizer_cn", None),
        organizer_email=getattr(t, "organizer_email", None),
        attendees=attendees,
        disable_forwarding=bool(getattr(t, "disable_forwarding", False) or False),
        calendar_priority=int(getattr(t, "calendar_priority", 2) or 2),
        calendar_categories=(getattr(t, "calendar_categories", None) or None),
        calendar_private=getattr(t, "calendar_private", None),
        meeting_begin=getattr(t, "meeting_begin", None),
        meeting_end=getattr(t, "meeting_end", None),
        meeting_provider=(getattr(t, "meeting_provider", None) or "none"),
        meeting_url=getattr(t, "meeting_url", None),
        meeting_tz=getattr(t, "meeting_tz", None),
        location=getattr(t, "location", None),
        allow_responses=bool(getattr(t, "allow_responses", True)),
        include_all_targets_as_attendees=bool(getattr(t, "include_all_targets_as_attendees", False)),
        is_calendar_reschedule=bool(getattr(t, "is_calendar_reschedule", False)),
        is_calendar_cancellation=bool(getattr(t, "is_calendar_cancellation", False)),
        meeting_uid=getattr(t, "meeting_uid", None),
        meeting_sequence=getattr(t, "meeting_sequence", None),
    )

@router.delete("/templates/{template_id}", tags=["Templates"])
def delete_template(template_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")

    active_statuses = ("running", "paused", "scheduled")
    active = (
        db.query(Campaign)
          .filter(Campaign.template_id == t.id, Campaign.status.in_(active_statuses))
          .all()
    )
    if active:
        names = ", ".join(c.name for c in active)
        raise HTTPException(
            status_code=409,
            detail=f"Template is used by active campaign(s): {names}. Stop/archive them before deleting."
        )

    # Snapshot missing
    db.query(Campaign).filter(
        Campaign.template_id == t.id,
        Campaign.template_name_snapshot.is_(None),
    ).update(
        {Campaign.template_name_snapshot: t.name},
        synchronize_session=False,
    )

    # Attachments will be removed (CASCADE)
    db.delete(t)
    db.commit()
    return {"ok": True}

@router.get("/templates/{template_id}/attachments", response_model=List[AttachmentOut], tags=["Templates"])
def list_template_attachments(template_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")
    
    rows = db.query(TemplateAttachment).filter(TemplateAttachment.template_id == template_id).all()

    return [AttachmentOut(id=r.id, filename=r.filename, mime_type=r.mime_type, size=len(r.data or b"")) for r in rows]

@router.post("/templates/{template_id}/attachments", response_model=List[AttachmentOut], tags=["Templates"])
def upload_template_attachments(template_id: int, files: List[UploadFile] = File(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")
    
    out: List[AttachmentOut] = []

    for f in files:
        data = f.file.read()
        filename = (f.filename or "file")

        # Replace existing attachments with the same filename (case insensitive)
        db.query(TemplateAttachment).filter(
            TemplateAttachment.template_id == template_id,
            func.lower(TemplateAttachment.filename) == func.lower(filename)
        ).delete(synchronize_session=False)

        att = TemplateAttachment(
            template_id=template_id,
            filename=filename,
            mime_type=f.content_type or "application/octet-stream",
            data=data,
        )

        db.add(att)
        db.flush()

        out.append(AttachmentOut(
            id=att.id,
            filename=att.filename,
            mime_type=att.mime_type,
            size=len(data)
        ))
    
    db.commit()
    return out

@router.delete("/templates/{template_id}/attachments/{attachment_id}", tags=["Templates"])
def delete_template_attachment(template_id: int, attachment_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    att = (
        db.query(TemplateAttachment)
          .filter(TemplateAttachment.id == attachment_id, TemplateAttachment.template_id == template_id)
          .first()
    )

    if not att:
        raise HTTPException(status_code=404, detail="Not found")
    
    db.delete(att)
    db.commit()
    return {"ok": True}

@router.get("/templates/{template_id}/calendar_attachments", response_model=List[AttachmentOut], tags=["Templates"])
def list_template_calendar_attachments(template_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")

    rows = (
        db.query(TemplateCalendarAttachment)
          .filter(TemplateCalendarAttachment.template_id == template_id)
          .all()
    )

    return [
        AttachmentOut(
            id=r.id, filename=r.filename, mime_type=r.mime_type, size=len(r.data or b"")
        )
        for r in rows
    ]

@router.post("/templates/{template_id}/calendar_attachments", response_model=List[AttachmentOut], tags=["Templates"])
def upload_template_calendar_attachments(template_id: int, files: List[UploadFile] = File(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(EmailTemplate).get(template_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")
    
    # Force not allowing calendar files on cancellation templates
    if getattr(t, "is_calendar_cancellation", False):
        raise HTTPException(status_code=422, detail="Calendar attachments are not allowed for cancellation templates")

    if not files:
        raise HTTPException(status_code=422, detail="No file provided")

    out: List[AttachmentOut] = []
    for f in files:
        data = f.file.read()
        filename = f.filename or "file"

        # Replace existing attachments with the same filename (case insensitive)
        db.query(TemplateCalendarAttachment).filter(
            TemplateCalendarAttachment.template_id == template_id,
            func.lower(TemplateCalendarAttachment.filename) == func.lower(filename)
        ).delete(synchronize_session=False)

        att = TemplateCalendarAttachment(
            template_id=template_id,
            filename=filename,
            mime_type=f.content_type or "text/calendar",
            data=data,
        )

        db.add(att)
        db.flush()

        out.append(AttachmentOut(
            id=att.id, filename=att.filename, mime_type=att.mime_type, size=len(data)
        ))

    db.commit()
    return out

@router.delete("/templates/{template_id}/calendar_attachments/{attachment_id}", tags=["Templates"])
def delete_template_calendar_attachment(template_id: int, attachment_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    att = (
        db.query(TemplateCalendarAttachment)
          .filter(
              TemplateCalendarAttachment.id == attachment_id,
              TemplateCalendarAttachment.template_id == template_id
          )
          .first()
    )

    if not att:
        raise HTTPException(status_code=404, detail="Not found")
    
    db.delete(att)
    db.commit()
    return {"ok": True}

@router.post("/templates/{template_id}/duplicate", response_model=TemplateOut, tags=["Templates"])
def duplicate_template(template_id: int, payload: TemplateDuplicate = Body(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    src = db.query(EmailTemplate).filter(EmailTemplate.id == template_id).first()

    if not src:
        raise HTTPException(status_code=404, detail="Template not found")

    exists_ = db.query(exists().where(EmailTemplate.name == payload.name)).scalar()
    if exists_:
        raise HTTPException(status_code=409, detail="A template with that name already exists")

    dst = EmailTemplate(
        name=payload.name,
        subject=src.subject,
        body_html=src.body_html,
        body_text=src.body_text,
        calendar_enabled=src.calendar_enabled,
        meeting_summary=src.meeting_summary,
        organizer_cn=src.organizer_cn,
        organizer_email=src.organizer_email,
        attendees=src.attendees,
        disable_forwarding=src.disable_forwarding,
        calendar_priority=src.calendar_priority,
        calendar_categories=src.calendar_categories,
        calendar_private=src.calendar_private,
        meeting_begin=src.meeting_begin,
        meeting_end=src.meeting_end,
        meeting_provider=src.meeting_provider or "none",
        meeting_url=src.meeting_url,
        meeting_tz=src.meeting_tz,
        location=src.location,
        allow_responses=src.allow_responses,
        is_calendar_reschedule=src.is_calendar_reschedule,
        is_calendar_cancellation=src.is_calendar_cancellation,
        meeting_uid=src.meeting_uid,
        meeting_sequence=src.meeting_sequence,
    )

    db.add(dst)
    db.flush()

    # Copy attachments
    atts = db.query(TemplateAttachment).filter(TemplateAttachment.template_id == src.id).all()
    for a in atts:
        db.add(TemplateAttachment(
            template_id=dst.id,
            filename=a.filename,
            mime_type=a.mime_type,
            data=a.data,
        ))

    # Copy calendar attachments only if not a cancellation
    cal_att = (
        db.query(TemplateCalendarAttachment)
          .filter(TemplateCalendarAttachment.template_id == src.id)
          .first()
    )
    if cal_att and not src.is_calendar_cancellation:
        db.add(TemplateCalendarAttachment(
            template_id=dst.id,
            filename=cal_att.filename,
            mime_type=cal_att.mime_type,
            data=cal_att.data,
        ))

    db.commit()
    db.refresh(dst)

    out_atts = [
        AttachmentOut(id=a.id, filename=a.filename, mime_type=a.mime_type, size=len(a.data or b""))
        for a in db.query(TemplateAttachment).filter(TemplateAttachment.template_id == dst.id).all()
    ]

    # Normalize attendees
    raw = getattr(dst, "attendees", None) or []
    att_list = []
    if isinstance(raw, list):
        for it in raw:
            cn = (it or {}).get("cn")
            em = (it or {}).get("email")
            if em:
                att_list.append(CalendarAttendee(cn=cn or "", email=em))

    return TemplateOut(
        id=dst.id,
        name=dst.name,
        subject=dst.subject,
        body_html=dst.body_html,
        body_text=dst.body_text,
        attachments=out_atts,
        calendar_enabled=dst.calendar_enabled,
        meeting_summary=dst.meeting_summary,
        organizer_cn=dst.organizer_cn,
        organizer_email=dst.organizer_email,
        attendees=att_list,
        disable_forwarding=dst.disable_forwarding,
        calendar_priority=dst.calendar_priority,
        calendar_categories=(dst.calendar_categories or None),
        calendar_private=dst.calendar_private,
        meeting_begin=dst.meeting_begin,
        meeting_end=dst.meeting_end,
        meeting_provider=(dst.meeting_provider or "none"),
        meeting_url=dst.meeting_url,
        meeting_tz=dst.meeting_tz,
        location=dst.location,
        allow_responses=dst.allow_responses,
         is_calendar_cancellation=bool(dst.is_calendar_cancellation),
        meeting_uid=dst.meeting_uid,
        meeting_sequence=dst.meeting_sequence,
    )

@router.post("/lists", response_model=TargetListOut, tags=["Target lists"])
def create_list(payload: TargetListCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if db.query(exists().where(TargetList.name == payload.name)).scalar():
        raise HTTPException(status_code=409, detail="A target list with that name already exists")
    
    tl = TargetList(name=payload.name)
    db.add(tl)

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A target list with that name already exists")

    db.refresh(tl)

    return TargetListOut(id=tl.id, name=tl.name)

@router.put("/lists/{list_id}", response_model=TargetListOut, tags=["Target lists"])
def rename_list(list_id: int, payload: TargetListUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    l = db.query(TargetList).get(list_id)

    if not l:
        raise HTTPException(status_code=404, detail="Not found")
    
    if payload.name != l.name and db.query(exists().where(TargetList.name == payload.name)).scalar():
        raise HTTPException(status_code=409, detail="A target list with that name already exists")
    
    l.name = payload.name

    try:
        db.commit()
    
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="A target list with that name already exists")

    db.refresh(l)

    return TargetListOut(id=l.id, name=l.name)

@router.get("/lists", response_model=List[TargetListOut], tags=["Target lists"])
def list_lists(db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = (
        db.query(
            TargetList.id,
            TargetList.name,
            func.count(Target.id).label("target_count"),
        )
        .outerjoin(Target, Target.list_id == TargetList.id)
        .group_by(TargetList.id)
        .order_by(TargetList.name.asc())
        .all()
    )

    return [
        TargetListOut(
            id=r.id,
            name=r.name,
            target_count=int(r.target_count or 0)
        )
        for r in rows
    ]

@router.post("/lists/{list_id}/targets", tags=["Target lists"])
def add_targets(list_id: int, targets: List[TargetCreate], db: Session = Depends(get_db), user=Depends(get_current_user)):
    for t in targets:
        db.add(Target(list_id=list_id, first_name=t.first_name, last_name=t.last_name, email=t.email, extra_field=t.extra_field, extra=t.extra))
    
    db.commit()
    return {"ok": True, "added": len(targets)}

@router.get("/lists/{list_id}/targets", tags=["Target lists"])
def list_targets(list_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(Target).filter(Target.list_id == list_id).all()
    return [{"id": r.id, "first_name": r.first_name, "last_name": r.last_name, "email": r.email, "extra_field": r.extra_field, "extra": r.extra} for r in rows]

@router.delete("/targets/{target_id}", tags=["Target lists"])
def delete_target(target_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    t = db.query(Target).get(target_id)

    if not t:
        raise HTTPException(status_code=404, detail="Not found")

    # Only block deletion if the target is already enqueued
    pending = (
        db.query(EmailLog)
          .filter(EmailLog.target_id == t.id, EmailLog.status == "pending")
          .count()
    )
    if pending > 0:
        raise HTTPException(
            status_code=409,
            detail="Target is pending in an active or scheduled campaign; stop/complete that campaign before deleting."
        )

    # Disassociate historical logs
    db.query(EmailLog).filter(EmailLog.target_id == t.id).update(
        { EmailLog.target_id: None }, synchronize_session=False
    )

    db.delete(t)
    db.commit()    

    return {"ok": True}

@router.delete("/lists/{list_id}", tags=["Target lists"])
def delete_list(list_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    l = db.get(TargetList, list_id)

    if not l:
        raise HTTPException(status_code=404, detail="Not found")

    active_refs = (
        db.query(Campaign.id, Campaign.name, Campaign.status)
          .join(CampaignTargetList, Campaign.id == CampaignTargetList.campaign_id)
          .filter(CampaignTargetList.list_id == list_id,
                  Campaign.status.notin_(["completed", "archived"]))
          .all()
    )

    if active_refs:
        sample = ", ".join(f"{n} ({s})" for _, n, s in active_refs[:3])
        more = "" if len(active_refs) <= 3 else f" (+{len(active_refs)-3} more)"

        raise HTTPException(
            status_code=409,
            detail=f"Cannot delete list: referenced by active/scheduled campaigns {sample}{more}."
        )

    # Disassociate historical logs
    tgt_ids_q = db.query(Target.id).filter(Target.list_id == list_id).subquery()
    db.query(EmailLog).filter(EmailLog.target_id.in_(tgt_ids_q)) \
        .update({EmailLog.target_id: None}, synchronize_session=False)

    # Unlink from campaigns
    db.query(CampaignTargetList).filter(CampaignTargetList.list_id == list_id) \
        .delete(synchronize_session=False)

    # Delete all targets in the list
    db.query(Target).filter(Target.list_id == list_id).delete(synchronize_session=False)

    db.delete(l)
    db.commit()

    return {"ok": True}

@router.post("/lists/{list_id}/duplicate", response_model=TargetListOut, status_code=201, tags=["Target lists"])
def duplicate_target_list(list_id: int, payload: TargetListDuplicate = Body(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    src = db.query(TargetList).get(list_id)

    if not src:
        raise HTTPException(status_code=404, detail="List not found")

    if db.query(exists().where(TargetList.name == payload.name)).scalar():
        raise HTTPException(status_code=409, detail="A target list with that name already exists")

    dst = TargetList(name=payload.name)
    db.add(dst)
    db.flush()

    # Copy all targets
    db.execute(
        text("""
            INSERT INTO targets (list_id, first_name, last_name, email, extra)
            SELECT :new_id, first_name, last_name, email, extra
            FROM targets
            WHERE list_id = :old_id
        """),
        {"new_id": dst.id, "old_id": src.id},
    )

    try:
        db.commit()
    
    except IntegrityError as e:
        db.rollback()
        raise HTTPException(status_code=409, detail="A target list with that name already exists")

    db.refresh(dst)

    return {"id": dst.id, "name": dst.name}

@router.get("/lists/search-by-email", tags=["Target lists"])
def search_lists_by_email(email: str, db: Session = Depends(get_db), user=Depends(get_current_user)):
    if not email or not email.strip():
        raise HTTPException(status_code=400, detail="Email is required")
    
    email_lower = email.strip().lower()
    
    # Find all lists that have targets with this email
    results = (
        db.query(
            TargetList.id,
            TargetList.name,
            func.count(Target.id).label("matching_targets")
        )
        .join(Target, TargetList.id == Target.list_id)
        .filter(func.lower(Target.email) == email_lower)
        .group_by(TargetList.id, TargetList.name)
        .order_by(TargetList.name.asc())
        .all()
    )
    
    return [
        {
            "id": r.id,
            "name": r.name,
            "matching_targets": r.matching_targets
        }
        for r in results
    ]

@router.post("/campaigns", response_model=CampaignOut, tags=["Campaigns"])
def create_campaign(payload: CampaignCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    # Normalize schedule
    sched = None
    if payload.scheduled_start_at:
        try:
            local_naive = datetime.fromisoformat(payload.scheduled_start_at)
            
            # Interpret as scheduler timezone
            tz = ZoneInfo(getattr(settings, "scheduler_timezone", "UTC"))
            if local_naive.tzinfo is None:
                local_aware = local_naive.replace(tzinfo=tz)
            else:
                local_aware = local_naive.astimezone(tz)
            
            # Convert to UTC for storage
            sched = local_aware.astimezone(ZoneInfo("UTC")).replace(tzinfo=None)
        except Exception:
            raise HTTPException(status_code=400, detail="scheduled_start_at must be an ISO datetime string")

    now = datetime.utcnow()
    status = "stopped"

    if payload.run_now:
        status = "running"
        sched = None
    
    elif sched and sched > now:
        status = "scheduled"

    # Prefill batch defaults if omitted
    if payload.use_batches:
        gs = {s.key: (s.value or "") for s in db.query(GlobalSetting).all()}

        if not payload.batch_size or payload.batch_size <= 0:
            try:
                payload.batch_size = int(gs.get("default_batch_size") or 0) or 60
            except Exception:
                payload.batch_size = 60

        if payload.batch_delay_seconds is None or payload.batch_delay_seconds < 0:
            try:
                payload.batch_delay_seconds = int(gs.get("default_batch_delay") or 0) or 60
            except Exception:
                payload.batch_delay_seconds = 60

    else:
        payload.batch_size = None
        payload.batch_delay_seconds = None

    names = [x.name for x in db.query(TargetList).filter(TargetList.id.in_(payload.list_ids)).all()]

    c = Campaign(
        name=payload.name,
        server_id=payload.server_id,
        profile_id=payload.profile_id,
        template_id=payload.template_id,
        selected_list_names=names,
        use_batches=payload.use_batches,
        batch_size=payload.batch_size,
        batch_delay_seconds=payload.batch_delay_seconds,
        custom_headers=payload.custom_headers or {},
        status=status,
        scheduled_start_at=sched,
        remove_duplicate_targets=payload.remove_duplicate_targets,
    )

    db.add(c)
    db.commit()
    db.refresh(c)

    # Insert campaign and list mappings
    for lid in payload.list_ids:
        db.add(CampaignTargetList(campaign_id=c.id, list_id=int(lid)))
    db.commit()

    # If run_now, enqueue mails
    if c.status == "running":
        template = db.query(EmailTemplate).get(c.template_id)
        _snapshot_calendar_fields_if_empty(db, c, template)

        existing = db.query(EmailLog).filter(EmailLog.campaign_id == c.id).count()
        if existing == 0:
            # DB Mutex for this campaign
            db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": int(c.id)})

            lids = [
                r.list_id
                for r in db.query(CampaignTargetList)
                        .filter(CampaignTargetList.campaign_id == c.id)
                        .all()
            ]

            if lids:
                if c.remove_duplicate_targets:
                    sub = (
                        db.query(func.min(Target.id).label("id"))
                          .filter(Target.list_id.in_(lids))
                          .group_by(func.lower(Target.email))
                    ).subquery()

                    base_q = db.query(Target).join(sub, Target.id == sub.c.id)
                
                else:
                    base_q = db.query(Target).filter(Target.list_id.in_(lids))

                # Skip emails already enqueued for this campaign
                already = db.query(EmailLog.id).filter(
                    EmailLog.campaign_id == c.id,
                    func.lower(EmailLog.to_email) == func.lower(Target.email)
                ).exists()

                targets = base_q.filter(~already).all()

                for t in targets:
                    db.add(EmailLog(
                        campaign_id=c.id,
                        target_id=t.id,
                        to_email=t.email,
                        status="pending",
                        target_first_name=t.first_name,
                        target_last_name=t.last_name,
                        target_extra_field=t.extra_field,
                    ))

                # Keep lock active until enqueue finishes
                db.commit()

        c.started_at = now
        c.scheduled_start_at = None
        db.commit()

        import threading
        from ..worker import run_campaign_send
        threading.Thread(target=run_campaign_send, args=(c.id,), daemon=True).start()

    return CampaignOut(id=c.id, name=c.name, status=c.status)

@router.get("/campaigns", tags=["Campaigns"])
def list_campaigns(db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(Campaign).order_by(Campaign.created_at.desc()).all()

    return [{
        "id": r.id,
        "name": r.name,
        "status": r.status,
        "created_at": _iso_utc(r.created_at),
        "scheduled_start_at": _iso_utc(r.scheduled_start_at)
    } for r in rows]

@router.get("/campaigns/{campaign_id}", response_model=CampaignDetailOut, tags=["Campaigns"])
def campaign_detail(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Not found")

    sent = db.query(EmailLog).filter(EmailLog.campaign_id==c.id, EmailLog.status=="sent").count()
    failed = db.query(EmailLog).filter(EmailLog.campaign_id==c.id, EmailLog.status=="failed").count()
    pending = db.query(EmailLog).filter(EmailLog.campaign_id==c.id, EmailLog.status=="pending").count()

    if (sent + failed + pending) > 0:
        total = sent + failed + pending

    else:
        lids = [r.list_id for r in db.query(CampaignTargetList).filter(CampaignTargetList.campaign_id==c.id).all()]
        total = db.query(Target).filter(Target.list_id.in_(lids)).count() if lids else 0

        if lids:
            if c.remove_duplicate_targets:
                # Count distinct emails across lists
                total = db.query(func.count(func.distinct(func.lower(Target.email)))).filter(Target.list_id.in_(lids)).scalar()
            else:
                total = db.query(func.count(Target.id)).filter(Target.list_id.in_(lids)).scalar()
        else:
            total = 0

    server = db.query(SendingServer).get(c.server_id) if c.server_id else None
    profile = db.query(SenderProfile).get(c.profile_id) if c.profile_id else None
    template = db.query(EmailTemplate).get(c.template_id)

    template = db.query(EmailTemplate).get(c.template_id)

    # Prefer stored snapshot on the campaign, fallback to template values
    meeting_summary = c.meeting_summary or (template.meeting_summary if template else None)
    meeting_begin = c.meeting_begin or (template.meeting_begin if template else None)
    meeting_end = c.meeting_end or (template.meeting_end if template else None)
    meeting_tz = c.meeting_tz or (template.meeting_tz if template else None) or settings.scheduler_timezone
    
    is_cancellation = bool(template and getattr(template, "is_calendar_cancellation", False))
    is_reschedule = bool(template and getattr(template, "is_calendar_reschedule", False))
    is_invite = bool(template and getattr(template, "calendar_enabled", False)) and not (is_cancellation or is_reschedule)

    # UI
    has_calendar_invite = is_invite
    has_calendar_reschedule = is_reschedule
    has_calendar_cancellation = is_cancellation

    gs = db.query(GlobalSetting).filter(GlobalSetting.key == "ics_prodid").first()
    resolved_prodid = (gs.value if gs and gs.value else settings.default_prodid)
    
    # Meeting provider & URL
    meeting_provider = (c.meeting_provider or (template.meeting_provider if template else None) or "none")
    meeting_url = c.meeting_url or (template.meeting_url if template else None)

    # Don't show URL for cancellations
    if has_calendar_cancellation:
        meeting_provider = "none"
        meeting_url = None

    return CampaignDetailOut(
        id=c.id, name=c.name, status=c.status,
        created_at=_iso_utc(c.created_at),
        started_at=_iso_utc(c.started_at),
        completed_at=_iso_utc(c.completed_at),
        scheduled_start_at=_iso_utc(c.scheduled_start_at),
        server=(server.name if server else (c.server_name_snapshot or "")),
        profile=(profile.profile_name if profile else (c.profile_name_snapshot or "")),
        profile_email=(profile.email if profile else (c.profile_email_snapshot or None)),
        template=(template.name if template else (c.template_name_snapshot or "")),
        lists=c.selected_list_names or [],
        total=total, sent=sent, failed=failed, pending=pending,
        custom_headers = (c.custom_headers or None),
        use_batches=bool(c.use_batches),
        batch_size=c.batch_size,
        batch_delay_seconds=c.batch_delay_seconds,
        remove_duplicate_targets=c.remove_duplicate_targets,
        last_error=c.last_error,
        last_error_trace=c.last_error_trace,
        last_error_at=_iso_utc(c.last_error_at) if c.last_error_at else None,
        has_calendar_invite=has_calendar_invite,
        has_calendar_reschedule=has_calendar_reschedule,
        has_calendar_cancellation=has_calendar_cancellation,
        calendar_uid=c.calendar_uid if (has_calendar_invite or has_calendar_reschedule or has_calendar_cancellation) else None,
        calendar_ics_prodid=resolved_prodid if (has_calendar_invite or has_calendar_reschedule or has_calendar_cancellation) else None,
        meeting_summary=meeting_summary,
        meeting_begin=meeting_begin,
        meeting_end=meeting_end,
        meeting_tz=meeting_tz,
        meeting_provider=meeting_provider,
        meeting_url=meeting_url,
        meeting_sequence=c.meeting_sequence if (has_calendar_reschedule or has_calendar_cancellation) else None,
    )

@router.get("/campaigns/{campaign_id}/targets_status", response_model=List[TargetStatus], tags=["Campaigns"])
def campaign_targets_status(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Not found")
    
    has_logs = db.query(EmailLog).filter(EmailLog.campaign_id==c.id).count() > 0

    # If the campaign has already enqueued any logs, show status and history
    if has_logs:
        rows = (
            db.query(EmailLog, Target)
              .outerjoin(Target, EmailLog.target_id == Target.id)
              .filter(EmailLog.campaign_id == c.id)
              .order_by(EmailLog.id.asc())
              .all()
        )

        out = []

        for log, t in rows:
            first_name = (log.target_first_name or (t.first_name if t else None))
            last_name = (log.target_last_name  or (t.last_name  if t else None))

            out.append(TargetStatus(
                id=(t.id if t else log.id), # If target is gone use log.id
                first_name=first_name,
                last_name=last_name,
                email=log.to_email,
                status=log.status,
                sent_at=_iso_utc(log.sent_at),
                error=log.error,
                opened=bool(getattr(log, "opened", False)),
                opened_at=_iso_utc(getattr(log, "opened_at", None)) if getattr(log, "opened_at", None) else None,
            ))
        return out
    
    # No logs yet, show preview
    lids = [r.list_id for r in db.query(CampaignTargetList).filter(CampaignTargetList.campaign_id==c.id).all()]
    out = []

    if lids:
        if c.remove_duplicate_targets:
            sub = (
                db.query(func.min(Target.id).label("id"))
                .filter(Target.list_id.in_(lids))
                .group_by(func.lower(Target.email))
            ).subquery()
            targets = db.query(Target).join(sub, Target.id == sub.c.id).all()
        
        else:
            targets = db.query(Target).filter(Target.list_id.in_(lids)).all()

        for t in targets:
            out.append(TargetStatus(
                id=t.id,
                first_name=t.first_name,
                last_name=t.last_name,
                email=t.email,
                status="not_enqueued",
                sent_at=None,
                error=None,
                opened=False,
                opened_at=None
            ))

    return out

@router.post("/campaigns/{campaign_id}/start", tags=["Campaigns"])
def start_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")

    if c.status == "error":
        c.last_error = None
        c.last_error_trace = None
        c.last_error_at = None
        c.status = "stopped"
        db.commit()

    if c.status in ("running","completed","archived"):
        raise HTTPException(status_code=400, detail=f"Cannot start campaign in status {c.status}")
    
    existing = db.query(EmailLog).filter(EmailLog.campaign_id==c.id).count()
    if existing == 0:
        # DB Mutex for this campaign
        db.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": int(c.id)})

        lids = [r.list_id for r in db.query(CampaignTargetList).filter(CampaignTargetList.campaign_id==c.id).all()]

        if lids:
            if c.remove_duplicate_targets:
                sub = (
                    db.query(func.min(Target.id).label("id"))
                    .filter(Target.list_id.in_(lids))
                    .group_by(func.lower(Target.email))
                ).subquery()

                base_q = db.query(Target).join(sub, Target.id == sub.c.id)

            else:
                base_q = db.query(Target).filter(Target.list_id.in_(lids))

            # Skip emails already enqueued for this campaign
            already = db.query(EmailLog.id).filter(
                EmailLog.campaign_id == c.id,
                func.lower(EmailLog.to_email) == func.lower(Target.email)
            ).exists()

            targets = base_q.filter(~already).all()

            for t in targets:
                db.add(EmailLog(
                    campaign_id=c.id,
                    target_id=t.id,
                    to_email=t.email,
                    status="pending",
                    target_first_name=t.first_name,
                    target_last_name=t.last_name,
                    target_extra_field=t.extra_field,
                ))

            # Keep lock active until enqueue finishes
            db.commit()

    template = db.query(EmailTemplate).get(c.template_id)
    _snapshot_calendar_fields_if_empty(db, c, template)
    _snapshot_used_vars_if_empty(db, c, template)

    c.status = "running"
    c.started_at = datetime.utcnow()
    c.scheduled_start_at = None
    db.commit()

    import threading
    from ..worker import run_campaign_send
    threading.Thread(target=run_campaign_send, args=(c.id,), daemon=True).start()

    return {"ok": True, "status": c.status}

@router.post("/campaigns/{campaign_id}/pause", tags=["Campaigns"])
def pause_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)
    
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if c.status != "running":
        raise HTTPException(status_code=400, detail="Campaign not running")
    
    c.status = "paused"
    db.commit()

    return {"ok": True, "status": c.status}

@router.post("/campaigns/{campaign_id}/resume", tags=["Campaigns"])
def resume_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c: 
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if c.status != "paused": 
        raise HTTPException(status_code=400, detail="Campaign not paused")
    
    # Check for failed/retrying emails and get the first one
    failed_log = (
        db.query(EmailLog)
        .filter(EmailLog.campaign_id == c.id)
        .filter(EmailLog.status.in_(["failed", "retrying"]))
        .order_by(EmailLog.id.asc())
        .first()
    )
    
    if failed_log:
        # Reset the failed/retrying email to pending
        print(f"[campaign {c.id}] Resetting failed/retrying email {failed_log.to_email} (ID: {failed_log.id}) to pending for retry", flush=True)
        failed_log.status = "pending"
        # Clear previous error
        failed_log.error = None
        db.commit()
    
    c.status = "running"
    db.commit()

    import threading
    from ..worker import run_campaign_send
    threading.Thread(target=run_campaign_send, args=(c.id,), daemon=True).start()

    return {"ok": True, "status": c.status}

@router.post("/campaigns/{campaign_id}/archive", tags=["Campaigns"])
def archive_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    # Clean up any pending emails before archiving
    _cleanup_pending_emails(db, c.id, "cancelled")

    c.status = "archived"
    c.completed_at = c.completed_at or datetime.utcnow()
    db.commit()

    return {"ok": True, "status": c.status}

@router.post("/campaigns/{campaign_id}/schedule", tags=["Campaigns"])
def schedule_campaign(campaign_id: int, payload: dict, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")

    raw = (payload or {}).get("scheduled_start_at")
    if not raw:
        raise HTTPException(status_code=400, detail="scheduled_start_at is required")

    # Convert to UTC for storage
    try:
        local_naive = datetime.fromisoformat(raw)
    except Exception:
        raise HTTPException(status_code=400, detail="scheduled_start_at must be an ISO datetime-local string")

    # Interpret as scheduler timezone
    tz = ZoneInfo(getattr(settings, "scheduler_timezone", "UTC"))
    if local_naive.tzinfo is None:
        local_aware = local_naive.replace(tzinfo=tz)
    else:
        local_aware = local_naive.astimezone(tz)

    # Convert to UTC for storage
    sched_utc_naive = local_aware.astimezone(ZoneInfo("UTC")).replace(tzinfo=None)

    # Reject past or now
    if sched_utc_naive <= datetime.utcnow():
        raise HTTPException(status_code=400, detail="Scheduled time must be in the future")

    if c.status in ("running", "completed", "archived"):
        raise HTTPException(status_code=400, detail=f"Cannot schedule campaign in status {c.status}")

    c.scheduled_start_at = sched_utc_naive
    c.status = "scheduled"
    db.commit()
    db.refresh(c)

    return {"ok": True, "status": c.status, "scheduled_start_at": c.scheduled_start_at.isoformat() + "Z"}

@router.post("/campaigns/{campaign_id}/unschedule", tags=["Campaigns"])
def unschedule_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if c.status == "scheduled":
        c.status = "stopped"
    
    c.scheduled_start_at = None
    db.commit()

    return {"ok": True, "status": c.status}

@router.get("/campaigns/{campaign_id}/logs", tags=["Campaigns"])
def campaign_logs(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(EmailLog).filter(EmailLog.campaign_id==campaign_id).order_by(EmailLog.id.asc()).all()
    
    return [{
        "id": r.id,
        "to_email": r.to_email,
        "status": r.status,
        "error": r.error,
        "sent_at": _iso_utc(r.sent_at)
    } for r in rows]

@router.get("/campaigns/{campaign_id}/variables", response_model=CampaignVariablesOut, tags=["Campaigns"])
def campaign_variables(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Not found")

    template = db.query(EmailTemplate).get(c.template_id) if c.template_id else None
    if not template and c.started_at is None:
        raise HTTPException(status_code=400, detail="Template not found")

    used = _campaign_used_vars(db, c, template)

    all_vars = db.query(GlobalVar).all()
    text_vars_used  = [v for v in all_vars if v.kind == "text"  and v.name in used]
    regex_vars_used = [v for v in all_vars if v.kind == "regex" and v.name in used]
    teams_vars_used = [v for v in all_vars if v.kind == "teams" and v.name in used]

    # Static vars: text + static teams meetings (only if used)
    static_vars: dict[str, str] = {v.name: (v.value or "") for v in text_vars_used}
    for v in teams_vars_used:
        if not v.per_recipient:
            static_vars[v.name] = v.value or ""

    # Regex + per-recipient teams meetings (only if used)
    headers = [v.name for v in sorted(regex_vars_used, key=lambda x: x.name.lower())]
    headers += [x.name for x in sorted(teams_vars_used, key=lambda x: x.name.lower()) if x.per_recipient]  # fixed scoping
    if "extraField" not in headers:
        headers = ["extraField", *headers]

    # If the template used randomID, include it as well
    if "randomID" in used and "randomID" not in headers:
        headers.append("randomID")

    # Join logs with targets
    q = (
        db.query(EmailLog, Target)
          .outerjoin(Target, EmailLog.target_id == Target.id)
          .filter(EmailLog.campaign_id == c.id)
          .order_by(EmailLog.id.asc())
          .all()
    )

    rows = []

    for log, t in q:
        if not getattr(log, "gen_vars", None):
            continue

        full_name = " ".join(x for x in [
            (log.target_first_name or (t.first_name if t else "") or ""),
            (log.target_last_name  or (t.last_name  if t else "") or "")
        ] if x).strip() or log.to_email

        # Build per-row values matching headers + extraField + randomID if present
        values = {h: (log.gen_vars or {}).get(h, "") for h in headers if h != "extraField"}
        values["extraField"] = (log.target_extra_field or "") or ((t.extra_field if t else "") or "")

        if "randomID" in headers:
            values["randomID"] = (log.gen_vars or {}).get("randomID", "")

        rows.append(CampaignVarRow(name=full_name, email=log.to_email, values=values))

    return {
        "static_vars": static_vars,
        "headers": headers,
        "rows": rows,
    }

@router.post("/campaigns/{campaign_id}/opened", response_model=MarkOpenedOut, tags=["Campaigns"])
def mark_opened(campaign_id: int, body: MarkOpenedIn, db: Session = Depends(get_db), user=Depends(get_user_both_auths)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Not found")
    
    if c.status == "archived":
        raise HTTPException(status_code=403, detail="Archived campaigns are read-only and cannot be edited")
    
    # Selectors
    by_email = bool(body.email)
    by_keyval = bool(body.var_name and body.var_value is not None)
    by_value_only = (body.var_value is not None and not body.var_name and not body.email)

    if sum([by_email, by_keyval, by_value_only]) != 1:
        raise HTTPException(status_code=400, detail="Provide only one selector: email, var_value, or var_name + var_value")

    bulk = bool(body.bulk) if body.bulk is not None else (True if by_value_only else False)
    q = db.query(EmailLog).filter(EmailLog.campaign_id == c.id)

    if by_email:
        q = q.filter(func.lower(EmailLog.to_email) == func.lower(body.email))

    elif by_keyval:
        q = q.filter(
            EmailLog.gen_vars.op('->>')(body.var_name) == str(body.var_value)
        )
    else:
        # Find value across any generated value
        q = q.filter(text(
            "EXISTS (SELECT 1 FROM json_each_text((email_logs.gen_vars)::json) kv "
            "        WHERE kv.value = :val)"
        ).bindparams(val=str(body.var_value)))

    if not bulk:
        q = q.order_by(EmailLog.id.asc()).limit(1)

    rows = q.all()

    if not rows:
        return MarkOpenedOut(updated=0)

    now = datetime.utcnow()
    target_state = bool(body.opened)
    changed = 0

    for r in rows:
        prev = bool(r.opened)

        if prev == target_state:
            if target_state and not r.opened_at:
                r.opened_at = now
                changed += 1
        else:
            r.opened = target_state
            r.opened_at = (now if target_state else None)
            changed += 1

    if changed:
        db.commit()

    return MarkOpenedOut(updated=changed)

@router.post("/campaigns/{campaign_id}/clear_error", tags=["Campaigns"])
def clear_campaign_error(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    c.last_error = None
    c.last_error_at = None

    if c.status == "error":
        c.status = "stopped"
    
    db.commit()
    return {"ok": True}

@router.get("/campaigns_calendar", tags=["Campaigns"])
def campaigns_calendar(db: Session = Depends(get_db), user=Depends(get_current_user)):
    items = []

    for c in db.query(Campaign).all():
        # Start time depends on campaign state
        if c.status == "scheduled" and c.scheduled_start_at:
            # Scheduled time
            start = c.scheduled_start_at
        elif c.started_at:
            # Started timestamp
            start = c.started_at
        else:
            # Show creation time
            start = c.created_at
        
        # End time
        end = c.completed_at or start
        
        items.append({
            "id": c.id,
            "title": c.name,
            "start": _iso_utc(start),
            "end": _iso_utc(end),
            "status": c.status
        })

    return items

@router.post("/campaigns/archive_all", tags=["Campaigns"])
def archive_all_campaigns(db: Session = Depends(get_db), user=Depends(get_current_user)):
    # Exclude running campaigns
    rows = (
        db.query(Campaign)
          .filter(Campaign.status != "archived")
          .filter(Campaign.status != "running")
          .all()
    )

    count = 0
    now = datetime.utcnow()

    for c in rows:
        c.status = "archived"
        c.archived_at = now if hasattr(c, "archived_at") else None
        c.scheduled_start_at = None
        count += 1

    db.commit()

    return {"ok": True, "archived": count}

@router.get("/user_campaigns", response_model=List[UserCampaignRow], tags=["Campaigns"])
def user_campaigns(email: EmailStr, db: Session = Depends(get_db), user=Depends(get_current_user)):
    e = email.strip().lower()

    rows = (
        db.query(
            Campaign.id.label("id"),
            Campaign.name.label("name"),
            func.coalesce(func.min(EmailLog.sent_at), Campaign.created_at).label("date"),
        )
        .join(EmailLog, EmailLog.campaign_id == Campaign.id)
        .filter(func.lower(EmailLog.to_email) == e)
        .group_by(Campaign.id, Campaign.name, Campaign.created_at)
        .order_by(func.coalesce(func.max(EmailLog.sent_at), Campaign.created_at).desc())
        .all()
    )

    return [UserCampaignRow(id=r.id, name=r.name, date=r.date) for r in rows]

@router.get("/campaigns/{campaign_id}/export", tags=["Campaigns"])
def export_campaign_results(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")

    logs = (
        db.query(EmailLog)
          .filter(EmailLog.campaign_id == c.id)
          .order_by(EmailLog.id.asc())
          .all()
    )

    # Figure out which variables were actually used by the template at the time
    template = db.query(EmailTemplate).get(c.template_id) if c.template_id else None
    used = _campaign_used_vars(db, c, template) if template else set()

    # Keys present in logs and used by the template
    union_keys: set[str] = set()
    for lg in logs:
        if lg.gen_vars:
            union_keys.update(lg.gen_vars.keys())

    var_headers = sorted(k for k in union_keys if k in used)

    # Ensure randomID is present if it was used
    if "randomID" in used and "randomID" not in var_headers:
        var_headers.append("randomID")

    # Excel generation
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = ["Name", "Email", "Extra field", "Time sent", "Mail status", "Opened", "Opened at", *var_headers]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    bold = Font(bold=True)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = bold
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    def full_name(lg: EmailLog) -> str:
        parts = [lg.target_first_name or "", lg.target_last_name or ""]
        return " ".join(p for p in parts if p).strip()

    for lg in logs:
        # Always add extraField
        target = db.query(Target).get(lg.target_id)
        extra_val = None

        if target is not None:
            if hasattr(target, "extraField"):
                extra_val = getattr(target, "extraField", None)
            else:
                extra_val = getattr(target, "extra_field", None)
        if extra_val is None:
            extra_val = lg.target_extra_field or ""

        opened = bool(getattr(lg, "opened", False))
        opened_at_dt = getattr(lg, "opened_at", None)
        opened_at = (
            opened_at_dt.isoformat(sep=" ", timespec="seconds")
            if (opened and opened_at_dt) else "-"
        )

        row = [
            full_name(lg) or lg.to_email,
            lg.to_email,
            (extra_val or ""),
            (lg.sent_at.isoformat(sep=" ", timespec="seconds") if lg.sent_at else ""),
            lg.status,
            opened,
            opened_at,
        ]
        for key in var_headers:
            row.append((lg.gen_vars or {}).get(key, ""))

        ws.append(row)

    # Autosize columns
    for col_cells in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(10, max_len + 2), 60)

    # Extra sheet for static vars
    static_vars_list: list[tuple[str, str]] = []
    if template:
        all_gvars = db.query(GlobalVar).all()
        for v in all_gvars:
            if v.name not in used:
                continue
            if v.kind == "text" or (v.kind == "teams" and not getattr(v, "per_recipient", False)):
                val = v.value or ""
                if not isinstance(val, str):
                    try:
                        val = json.dumps(val, ensure_ascii=False)
                    except Exception:
                        val = str(val)
                static_vars_list.append((v.name, val))

    if static_vars_list:
        ws2 = wb.create_sheet(title="Static variables")
        ws2.append(["Variable", "Value"])

        for col in range(1, 3):
            cell = ws2.cell(row=1, column=col)
            cell.font = bold
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for name, value in static_vars_list:
            ws2.append([name, value])

        for col_cells in ws2.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
            ws2.column_dimensions[col_cells[0].column_letter].width = min(max(10, max_len + 2), 60)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"{c.name}_Results.xlsx"

    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.delete("/campaigns/{campaign_id}", tags=["Campaigns"])
def delete_campaign(campaign_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    c = db.query(Campaign).get(campaign_id)

    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if c.status != "archived":
        raise HTTPException(status_code=400, detail="Only archived campaigns can be deleted")
    
    # Remove dependent rows
    db.query(EmailLog).filter(EmailLog.campaign_id == c.id).delete(synchronize_session=False)
    db.query(CampaignTargetList).filter(CampaignTargetList.campaign_id == c.id).delete(synchronize_session=False)
    
    db.commit()
    db.delete(c)
    db.commit()

    return {"ok": True}

@router.post("/campaigns/delete_all", tags=["Campaigns"])
def delete_all_archived_campaigns(db: Session = Depends(get_db), user=Depends(get_current_user)):
    rows = db.query(Campaign).filter(Campaign.status == "archived").all()
    count = 0

    for c in rows:
        db.query(EmailLog).filter(EmailLog.campaign_id == c.id).delete(synchronize_session=False)
        db.query(CampaignTargetList).filter(CampaignTargetList.campaign_id == c.id).delete(synchronize_session=False)
        db.delete(c)
        count += 1
    
    db.commit()
    
    return {"ok": True, "deleted": count}